from flask import Flask, render_template, request, redirect, url_for, flash, send_file, session
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
import os
from openpyxl import load_workbook
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path
import sqlite3, json, shutil, subprocess, re, io
from datetime import datetime, date

BASE=Path(__file__).resolve().parent
DB=BASE/'contratos.db'; UP=Path(os.getenv('UPLOAD_DIR', str(BASE/'uploads'))); OUT=Path(os.getenv('OUTPUT_DIR', str(BASE/'gerados')))
UP.mkdir(parents=True,exist_ok=True); OUT.mkdir(parents=True,exist_ok=True)
DATABASE_URL=os.getenv('DATABASE_URL','').strip()
app=Flask(__name__); app.secret_key=os.getenv('SECRET_KEY','troque-esta-chave-em-producao'); ACTION_PASSWORD=os.getenv('ACTION_PASSWORD','123456'); STATUS_OPTS=['PENDENTE','EM ELABORAÇÃO','CONTRATO GERADO','AGUARDANDO ASSINATURA','ATIVO','CANCELADO']
app.config['MAX_CONTENT_LENGTH']=12*1024*1024

FIELDS=['razao_social','nome_fantasia','endereco','numero','bairro','cep','cidade','uf','emite_mtr','link_maps','lat_long','cpf_cnpj','ie','telefone_empresa','atividade','contato','email_nfe','telefone_contato','email_contato','responsavel','email_responsavel','cpf_responsavel','rg','cargo','representante','faturado_por','coletado_por','inicio','limpeza_csao','periodicidade','franquia_agua','franquia_solidos','franquia_terra','valor_faturamento','valor_agua','valor_excedente','valor_terra','valor_lampada','outros_valores','deslocamento','faturamento_minimo','dia_limite_faturamento','categoria','faturamento_mensal','medicao_diaria','dia_especifico','tipo_prazo','parcela','prazo_dia_fixo','faturamento_terceiro','cliente_gerador','observacoes']
for i in range(1,6): FIELDS += [f'kit_emb_{i}',f'kit_cor_{i}',f'kit_qtd_{i}']

def _is_pg(): return DATABASE_URL.startswith('postgresql://') or DATABASE_URL.startswith('postgres://')

class DBConn:
 def __init__(self):
  self.pg=_is_pg()
  if self.pg:
   import psycopg
   from psycopg.rows import dict_row
   url=DATABASE_URL.replace('postgres://','postgresql://',1)
   self.c=psycopg.connect(url,row_factory=dict_row)
  else:
   self.c=sqlite3.connect(DB); self.c.row_factory=sqlite3.Row
 def execute(self,sql,params=()):
  if self.pg: sql=sql.replace('?', '%s')
  return self.c.execute(sql,params)
 def commit(self): self.c.commit()
 def rollback(self): self.c.rollback()
 def close(self): self.c.close()

def con(): return DBConn()

def init():
 c=con(); cols=', '.join([f'"{x}" TEXT' for x in FIELDS])
 if _is_pg():
  c.execute(f"CREATE TABLE IF NOT EXISTS contratos(id SERIAL PRIMARY KEY,chave TEXT,status TEXT DEFAULT 'PENDENTE',criado_em TEXT,{cols})")
  c.execute("CREATE TABLE IF NOT EXISTS usuarios(id SERIAL PRIMARY KEY, usuario TEXT UNIQUE NOT NULL, nome TEXT, senha_hash TEXT NOT NULL, perfil TEXT NOT NULL DEFAULT 'USUARIO', ativo INTEGER NOT NULL DEFAULT 1)")
 else:
  c.execute(f"CREATE TABLE IF NOT EXISTS contratos(id INTEGER PRIMARY KEY AUTOINCREMENT,chave TEXT,status TEXT DEFAULT 'PENDENTE',criado_em TEXT,{cols})")
  c.execute("CREATE TABLE IF NOT EXISTS usuarios(id INTEGER PRIMARY KEY AUTOINCREMENT, usuario TEXT UNIQUE NOT NULL, nome TEXT, senha_hash TEXT NOT NULL, perfil TEXT NOT NULL DEFAULT 'USUARIO', ativo INTEGER NOT NULL DEFAULT 1)")
 admin=os.getenv('ADMIN_USER','admin'); pwd=os.getenv('ADMIN_PASSWORD','123456')
 existing_admin=c.execute('select id from usuarios where usuario=?',(admin,)).fetchone()
 if not existing_admin:
  c.execute('insert into usuarios(usuario,nome,senha_hash,perfil,ativo) values(?,?,?,?,1)',(admin,'Administrador',generate_password_hash(pwd),'ADMIN'))
 else:
  # Mantém a senha do administrador sincronizada com ADMIN_PASSWORD do Render.
  c.execute('update usuarios set senha_hash=?, perfil=?, ativo=1 where usuario=?',(generate_password_hash(pwd),'ADMIN',admin))
 c.commit(); c.close()
init()

def login_required(fn):
 @wraps(fn)
 def wrapper(*args,**kwargs):
  if not session.get('user_id'): return redirect(url_for('login',next=request.path))
  return fn(*args,**kwargs)
 return wrapper

def admin_required(fn):
 @wraps(fn)
 def wrapper(*args,**kwargs):
  if not session.get('user_id'): return redirect(url_for('login'))
  if session.get('perfil')!='ADMIN': flash('Acesso exclusivo do administrador.'); return redirect(url_for('index'))
  return fn(*args,**kwargs)
 return wrapper

def sval(v):
 if v is None:return ''
 if isinstance(v,(datetime,date)):return v.strftime('%d/%m/%Y')
 return str(v).strip()

def read_capa(path):
 wb=load_workbook(path,data_only=True); ws=wb['CAPA']
 # Valores ficam nas células imediatamente à direita/área mesclada do rótulo.
 m={'razao_social':'C7','nome_fantasia':'C8','endereco':'C9','numero':'J9','bairro':'C10','cep':'J10','cidade':'C11','uf':'F11','emite_mtr':'I11','link_maps':'C12','lat_long':'I12','cpf_cnpj':'C13','telefone_empresa':'H13','ie':'C14','atividade':'H14','contato':'C15','email_nfe':'H15','telefone_contato':'C16','email_contato':'H16','responsavel':'C17','email_responsavel':'H17','cpf_responsavel':'C18','rg':'H18','cargo':'C19','representante':'H19','faturado_por':'D21','coletado_por':'I21','inicio':'D22','limpeza_csao':'I22','periodicidade':'D23','franquia_agua':'I23','franquia_solidos':'D24','franquia_terra':'I24','valor_faturamento':'D26','valor_agua':'I26','valor_excedente':'D27','valor_terra':'I27','valor_lampada':'D28','outros_valores':'I28','deslocamento':'D29','faturamento_minimo':'H29','dia_limite_faturamento':'J29','categoria':'D30','faturamento_mensal':'D30','medicao_diaria':'H30','dia_especifico':'J30','tipo_prazo':'D31','parcela':'H31','prazo_dia_fixo':'J31','faturamento_terceiro':'D32','cliente_gerador':'I32','observacoes':'B40'}
 d={k:sval(ws[v].value) for k,v in m.items()}
 for i,row in enumerate(range(34,39),1):
  d[f'kit_emb_{i}']=sval(ws[f'C{row}'].value); d[f'kit_cor_{i}']=sval(ws[f'F{row}'].value); d[f'kit_qtd_{i}']=sval(ws[f'J{row}'].value)
 # Listas reais da aba SELECOES
 s=wb['SELECOES']
 def vals(col,rows): return [s.cell(r,col).value for r in rows if s.cell(r,col).value not in (None,'')]
 o={'pagamento':vals(1,range(3,6)),'filial':vals(3,range(3,6)),'representante':vals(4,range(3,6)),'mtr':vals(5,range(3,6)),'cond':vals(6,range(3,6)),'cor':vals(1,range(13,17)),'periodicidade':vals(2,range(12,17)),'embalagem':vals(4,range(12,19)),'categoria':vals(5,range(12,18)),'faturamento':vals(1,range(40,46))}
 # Formata campos monetários importados no padrão brasileiro.
 for k in ['valor_faturamento','valor_agua','valor_excedente','valor_terra','valor_lampada','deslocamento','faturamento_minimo']:
  if d.get(k)!='': d[k]=fmt_br(d[k],2)
 return d,o

def parse_num_br(v):
 if v in (None,''): return None
 if isinstance(v,(int,float)): return float(v)
 s=str(v).strip().replace('R$','').replace(' ','')
 try:
  if ',' in s:
   s=s.replace('.','').replace(',','.')
  return float(s)
 except: return None

def fmt_br(v, casas=2):
 n=parse_num_br(v)
 if n is None: return ''
 x=f'{n:,.{casas}f}'
 return x.replace(',','X').replace('.',',').replace('X','.')

def money(v):
 x=fmt_br(v,2)
 return f'R$ {x}' if x else 'R$ 0,00'

def filial_key(nome):
 n=(nome or '').upper().strip()
 if 'NORTE' in n or 'ECONORTE' in n or 'SINOP' in n: return 'NORTE'
 if 'CAMPO' in n or 'GRANDE' in n: return 'CAMPO_GRANDE'
 return 'CUIABA'

FILIAIS={
 'CAMPO_GRANDE':{
  'nome':'ECOSUPPLY RECICLADORA LTDA','cidade':'Campo Grande','uf':'MS','estado':'Mato Grosso do Sul',
  'endereco':'Av. Consul Assaf Trad, 2.739, Coronel Antonino, CEP 79.013-545',
  'cnpj':'10.533.843/0001-07','ie':'28.350.688-1','representante':'Valdiney Pedro Rodrigues','cargo':'Gerente Geral','cpf':'051.190.046-52',
  'cabecalho':['ECOSUPPLY RECICLADORA LTDA.','Av. Cônsul Assaf Trad, 2739','Coronel Antonino','Campo Grande – MS','CEP: 79013-545','TEL: (67) 3373-0104','www.supplyservice.com.br'],
  'foro':'Campo Grande/MS'
 },
 'CUIABA':{
  'nome':'ECOSUPPLY RECICLADORA LTDA','cidade':'Cuiabá','uf':'MT','estado':'Mato Grosso',
  'endereco':'Rua Paulo Masayuki Uezato, S/N, Quadra ind. 7 Lote 22 a 24, Distrito Industrial, CEP 78.098-400',
  'cnpj':'10.533.843/0002-98','ie':'13.432.014-0','representante':'Valdiney Pedro Rodrigues','cargo':'Gerente Geral','cpf':'051.190.046-52',
  'cabecalho':['ECOSUPPLY RECICLADORA LTDA','Rua Paulo Masayuki Uezato, S/N, Quadra ind. 7 Lote 22 a 24','Distrito Industrial','Cuiabá – MT','CEP 78.098-400','www.supplyservice.com.br'],
  'foro':'Campo Grande/MS'
 },
 'NORTE':{
  'nome':'ECOSUPPLY RECICLADORA DO NORTE LTDA','cidade':'Sinop','uf':'MT','estado':'Mato Grosso',
  'endereco':'Estrada Municipal Castanheira, S/N, Lote 01 Quadra 11, Bairro Lídia',
  'cnpj':'62.463.190/0001-30','ie':'14.145.716-3','representante':'','cargo':'','cpf':'',
  'cabecalho':['ECOSUPPLY RECICLADORA DO NORTE LTDA','Estrada Municipal Castanheira, S/N, Lote 01 Quadra 11','Bairro Lídia','Sinop - MT','CNPJ: 62.463.190/0001-30','CEP: 78559-899'],
  'foro':'Sinop/MT'
 }
}

def filial_data(nome): return FILIAIS[filial_key(nome)]

def filial_text(nome):
 f=filial_data(nome)
 txt=f"{f['nome']}, com sede no Município de {f['cidade']}, Estado de {f['estado']}, na {f['endereco']}, devidamente inscrita no CNPJ sob nº {f['cnpj']} e Inscrição Estadual sob nº {f['ie']}"
 if f['representante']:
  txt += f", neste ato representada pelo seu bastante procurador {f['representante']}, ocupante do cargo {f['cargo']} e portador do CPF: {f['cpf']}"
 return txt + ', doravante denominada CONTRATADA.'

def addp(doc,text,bold=False,center=False):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after=Pt(5); r=p.add_run(text); r.bold=bold; r.font.name='Arial'; r.font.size=Pt(10); return p

def build_doc(r,path):
 d=dict(r); doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
 style=doc.styles['Normal']; style.font.name='Arial'; style._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); style.font.size=Pt(10)
 addp(doc,'INSTRUMENTO PARTICULAR DE CONTRATO DE PRESTAÇÃO DE SERVIÇOS DE RETIRADA, DECOMPOSIÇÃO, TRATAMENTO E DESCARTE',True,True)
 addp(doc,'IDENTIFICAÇÃO DAS PARTES CONTRATANTES',True,True)
 ident=f"{d.get('razao_social','')}, pessoa jurídica/pessoa física com sede em {d.get('cidade','')}, Estado de {d.get('uf','')}, na {d.get('endereco','')}, nº {d.get('numero','')} - Bairro {d.get('bairro','')} - CEP: {d.get('cep','')}, inscrita no CPF/CNPJ sob n.º {d.get('cpf_cnpj','')} e Inscrição Estadual n.º {d.get('ie','')}, neste ato representada por {d.get('responsavel','')}, RG n.º {d.get('rg','')} e CPF n.º {d.get('cpf_responsavel','')}, doravante denominada CONTRATANTE."
 addp(doc,ident); addp(doc,filial_text(d.get('faturado_por')))
 addp(doc,'As partes acima identificadas têm, entre si, justo e acertado o presente Contrato de Prestação de Serviços, que se regerá pelas cláusulas seguintes e pelas condições de preço, forma e termo de pagamento descritas no presente.')
 addp(doc,'DO OBJETO DO CONTRATO',True,True)
 residuos=d.get('observacoes') or 'resíduos sólidos classe I e classe II conforme cadastro e caracterização fornecida pela CONTRATANTE'
 kits=[]
 for i in range(1,6):
  if d.get(f'kit_emb_{i}') and d.get(f'kit_qtd_{i}'): kits.append(f"- {d[f'kit_qtd_{i}']} unidade(s) de {d[f'kit_emb_{i}']} - cor {d.get(f'kit_cor_{i}','')};")
 kittext='\n'.join(kits) if kits else '- Não há kit de coleta/comodato informado.'
 addp(doc,f"Cláusula 1ª. É objeto do presente contrato a prestação dos serviços de:\nA) Serviço de coleta, triagem e envio para tratamento/destinação de {residuos}.\nB) Fornecimento a título de comodato, sem custo adicional, de:\n{kittext}\nC) Transporte dos resíduos em caminhão licenciado para transporte de carga perigosa e com motorista habilitado com MOPP.\nD) Destinação ambientalmente adequada de todos os resíduos coletados.\nParágrafo único – Caberá à CONTRATANTE providenciar às suas expensas, quando se fizer necessário, a caracterização dos resíduos de que trata esta Cláusula, através de análises completas executadas em entidades ou laboratórios de reconhecida capacidade técnica, conforme técnicas e métodos estabelecidos pela Norma Técnica, Legal ou Regulamentar pertinente, fornecendo, à CONTRATADA, cópia do laudo assinado por profissional competente.")
 addp(doc,'OBRIGAÇÕES DA CONTRATANTE',True,True)
 addp(doc,'Cláusula 2ª. A CONTRATANTE deverá abster-se de transferir para os locais de coleta quaisquer dos itens enumerados na cláusula primeira originados de outros estabelecimentos geradores, ainda que da mesma empresa ou grupo econômico.')
 mtr=d.get('emite_mtr','')
 mtrtxt='Compete ao CONTRATANTE/GERADOR a responsabilidade pela emissão do Manifesto de Transporte de Resíduos – MTR, por meio do SINIR, para cada remessa de resíduos destinada à coleta e/ou destinação.' if 'CLIENTE' in mtr.upper() else 'A CONTRATANTE autoriza a operacionalização da emissão dos MTRs pela CONTRATADA, devendo cadastrá-la no SINIR com as permissões necessárias, permanecendo com a CONTRATANTE as demais obrigações ambientais vinculadas à geração de resíduos e declarações aos órgãos reguladores.'
 addp(doc,'Parágrafo Primeiro – Para fins deste instrumento, consideram-se “geradores” as pessoas físicas ou jurídicas que, em decorrência de sua atividade, gerem os resíduos objeto deste contrato.\nParágrafo Segundo – '+mtrtxt)
 addp(doc,'Cláusula 3ª. A CONTRATANTE deverá fornecer condições mínimas para atender a Segurança do Trabalho ao colaborador da Supply no momento da efetivação dos serviços objetos do presente instrumento, compreendendo todas as Normas Regulamentadoras do Ministério do Trabalho e pela legislação aplicável.

Parágrafo Único: A CONTRATANTE deverá atender, integralmente, a Norma Regulamentadora de n° 16, estando condicionada a prestação dos serviços ao fato de que toda a área de operação deverá abranger, no mínimo, círculo de distância do colaborador com raio de 10 metros com centro no ponto de abastecimento.')
 addp(doc,'Cláusula 4ª. A CONTRATANTE deverá armazenar os itens discriminados na cláusula primeira em instalações adequadas devidamente licenciadas pelo órgão ambiental competente, em lugar acessível à coleta, utilizando os recipientes fornecidos a título de comodato pela CONTRATADA, de modo a não contaminar o meio ambiente.')
 addp(doc,'OBRIGAÇÕES DA CONTRATADA',True,True)
 clauses={5:'É dever da CONTRATADA observar todas as normas ambientais e técnicas pertinentes à execução dos serviços, dando plena e total garantia dos mesmos e responsabilizando-se isoladamente pela qualidade dos serviços prestados, devendo reparar ou refazer qualquer serviço executado em desconformidade com as especificações, instruções e normas técnicas.',6:'É dever da CONTRATADA estar devidamente autorizada e licenciada pelos órgãos competentes para realizar as atividades objeto deste contrato.'}
 for n,t in clauses.items(): addp(doc,f'Cláusula {n}ª. {t}')
 per=d.get('periodicidade') or 'CONFORME PROGRAMAÇÃO'; fran=d.get('franquia_solidos') or '0'
 addp(doc,f"Cláusula 7ª. A CONTRATADA deverá efetuar a coleta e destinação na periodicidade {per}, em quantidade não superior a {fran} kg de todo material a ser coletado, sendo que a eventual necessidade de coleta fora da periodicidade programada ou em quantidade superior implicará no pagamento de taxa adicional, conforme condições deste instrumento.")
 for n,t in [(8,'A CONTRATADA se reserva ao direito de devolver à CONTRATANTE, sem ônus para a primeira, todo o resíduo recebido que não possa processá-lo em razão de impedimento legal ou por inviabilidade técnica ou econômica.'),(9,'A CONTRATADA deverá manter registro das ocorrências que não se enquadrem como rotina.'),(10,'A CONTRATADA deverá comunicar à CONTRATANTE, por escrito ou por meio eletrônico, quaisquer ocorrências que impeçam o cumprimento das obrigações por ela assumidas, indicando suas causas, efeitos e sugestões que devam ser tomadas.'),(11,'As partes se obrigam a obedecer às exigências, condições e determinações regulares do Órgão de Controle da Poluição Ambiental competente, responsabilizando-se, cada uma, isoladamente, pelas infrações que eventualmente cometerem.')]: addp(doc,f'Cláusula {n}ª. {t}')
 addp(doc,'Cláusula 12ª. Para a execução dos serviços, a CONTRATADA disponibilizará os recursos necessários, incluindo veículo devidamente certificado e licenciado, condutores habilitados e treinados, mão de obra capacitada, EPI’s e documentação aplicável.')
 addp(doc,'DA COLETA E DO TRANSPORTE',True,True)
 addp(doc,'Cláusula 13ª. A CONTRATADA irá realizar as coletas no período diurno, no estabelecimento da CONTRATANTE, não sendo executadas aos sábados, domingos e feriados. Caso o dia programado para coleta seja feriado, a mesma será realizada no dia útil anterior ou posterior àquela data.')
 addp(doc,'Cláusula 14ª. A CONTRATANTE se responsabilizará por acompanhar a coleta, a pesagem e o preenchimento do manifesto, dando ciência de todos os resíduos que foram coletados.')
 addp(doc,'Cláusula 15ª. Ocorrendo impossibilidade real da CONTRATADA na execução do serviço, esta deverá ser reprogramada e comunicada imediatamente à CONTRATANTE para a sua realização.')
 addp(doc,'DO PREÇO E DAS CONDIÇÕES DE PAGAMENTO',True,True)
 preco=[f"- {money(d.get('valor_faturamento'))} referente ao serviço na periodicidade {per}."]
 for key,label,un in [('valor_excedente','material coletado excedente','kg'),('valor_terra','terra, serragem e graxa','kg'),('valor_agua','água coletada','litro'),('valor_lampada','lâmpada coletada','unidade')]:
  if d.get(key): preco.append(f"- {money(d[key])} por {un} de {label}.")
 if d.get('deslocamento'): preco.append(f"- {money(d['deslocamento'])} referente a deslocamento, quando aplicável.")
 if d.get('outros_valores'): preco.append(f"- Outros valores: {d['outros_valores']}.")
 pag=f"Tipo de prazo: {d.get('tipo_prazo','')}; prazo/dia fixo: {d.get('prazo_dia_fixo','')}; parcela: {d.get('parcela','')}; dia limite de faturamento: {d.get('dia_limite_faturamento','')}."
 addp(doc,'Cláusula 16ª. A CONTRATANTE pagará à CONTRATADA, mediante cobrança, a título de remuneração pelos serviços avençados por este instrumento, os seguintes valores:\n'+'\n'.join(preco)+'\n'+pag)
 addp(doc,'Parágrafo Primeiro: O atraso no pagamento de toda e qualquer fatura obrigará a CONTRATANTE ao pagamento do principal acrescido de multa contratual de 2% (dois por cento), além de juros de 1,0% (um por cento) ao mês, sobre o total do débito.\nParágrafo Segundo: A remuneração inclui o tratamento e a disposição final dos resíduos resultantes do processamento, de modo a atender às exigências dos órgãos ambientais competentes.\nParágrafo Terceiro: Eventuais tributos e contribuições que venham a ser criados ou elevados serão objeto de repasse para os custos.\nParágrafo Quarto: A CONTRATADA emitirá Nota Fiscal Eletrônica e a remeterá para o e-mail '+(d.get('email_nfe') or '[E-MAIL NF-e]')+'.\nParágrafo Quinto: Se o boleto bancário não puder ser pago e a CONTRATADA autorizar pagamento por depósito, poderão ser aplicadas as despesas bancárias pertinentes.\nParágrafo Sexto: O valor da remuneração será reajustado automaticamente a cada 12 (doze) meses através da aplicação do índice IGP-M ou outro índice ou reajuste pertinente.')
 addp(doc,'Cláusula 17ª. No que se refere à incidência do Imposto Sobre Serviços - ISS, a CONTRATANTE deverá observar rigorosamente os dispositivos legais que tratam da incidência e eventual retenção, sendo que a exigência ou retenção indevida do tributo implicará no repasse do valor para o preço do serviço.')
 addp(doc,'DA RESCISÃO',True,True)
 for n,t in [(18,'No caso de descumprimento de qualquer das cláusulas deste instrumento, a parte que não cumpriu deverá pagar multa correspondente ao valor dos serviços nos últimos três meses anteriores à infração contratual, além de ensejar a rescisão contratual motivada.'),(19,'O descumprimento de qualquer das cláusulas deste instrumento permite à parte lesada a rescisão contratual motivada, bastando para tanto a notificação da parte contrária, que produzirá efeitos imediatos.'),(20,'Poderá o presente instrumento ser rescindido por qualquer uma das partes, em qualquer momento, sem motivo relevante, devendo a outra parte ser avisada previamente por escrito, no prazo de 30 (trinta) dias.'),(21,'Caso a CONTRATANTE já tenha realizado o pagamento pelo serviço e requisite a rescisão imotivada, terá o valor da quantia paga devolvido, deduzindo-se 10% de taxas administrativas e impostos recolhidos.'),(22,'Caso seja a CONTRATADA quem requeira a rescisão imotivada, deverá devolver a quantia referente aos serviços não prestados à CONTRATANTE, acrescida de 2% de taxas administrativas.')]: addp(doc,f'Cláusula {n}ª. {t}')
 addp(doc,'DO PRAZO',True,True); addp(doc,'Cláusula 23ª. O presente contrato é por prazo indeterminado iniciando-se no momento e na data da assinatura do mesmo.')
 addp(doc,'DAS CONDIÇÕES GERAIS',True,True); addp(doc,'Cláusula 24ª. Salvo com a expressa autorização da CONTRATANTE, não pode a CONTRATADA transferir ou subcontratar os serviços previstos neste instrumento, sob o risco de ocorrer a rescisão imediata.'); addp(doc,'Cláusula 25ª. A CONTRATADA emitirá o Certificado de Tratamento e Disposição Final de Resíduos, no qual constará a quantidade tratada no período correspondente, sendo enviado ao cliente no prazo de 90 dias.')
 addp(doc,'DO FORO',True,True); addp(doc,'Cláusula 26ª. Para dirimir quaisquer controvérsias oriundas do presente contrato, as partes elegem o foro da comarca de Campo Grande/MS.')
 local=d.get('cidade') or 'Cuiabá'; addp(doc,f"{local}, {datetime.now().strftime('%d/%m/%Y')}.",False,True); addp(doc,'Por estarem assim justos e contratados, firmam o presente instrumento em 02 (duas) vias de igual teor, juntamente com 02 (duas) testemunhas.',False,True)
 addp(doc,'\n________________________________________\nECOSUPPLY RECICLADORA LTDA\nCONTRATADA',False,True); addp(doc,f"\n________________________________________\n{d.get('razao_social','')}\nCONTRATANTE",False,True)
 doc.save(path)


def authorized(id):
    return id in session.get('authorized_contracts', [])

def protect(id):
    if not authorized(id):
        flash('Informe a senha para acessar as ações deste contrato.')
        return redirect(url_for('acesso', id=id))
    return None

def _set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
 tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
 if tcMar is None:
  tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
 for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  node=tcMar.find(qn('w:'+m))
  if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
  node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def _set_bottom_border(paragraph):
 pPr=paragraph._p.get_or_add_pPr(); pBdr=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
 bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'18'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),'000000')
 pBdr.append(bottom); pPr.append(pBdr)

def doc_header(section, filial=''):
 f=filial_data(filial); header=section.header
 header.is_linked_to_previous=False
 t=header.add_table(rows=1, cols=2, width=Cm(17.6)); t.autofit=False
 t.columns[0].width=Cm(10.4); t.columns[1].width=Cm(7.2)
 for c in t.rows[0].cells: c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; _set_cell_margins(c,0,0,0,0)
 p=t.cell(0,0).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 logo=BASE/'static'/'logo_ecosupply.png'
 if logo.exists(): p.add_run().add_picture(str(logo), width=Cm(6.5))
 p2=t.cell(0,1).paragraphs[0]; p2.alignment=WD_ALIGN_PARAGRAPH.LEFT; p2.paragraph_format.space_after=Pt(0); p2.paragraph_format.line_spacing=0.92
 for i,line in enumerate(f['cabecalho']):
  rr=p2.add_run(line + ('\n' if i < len(f['cabecalho'])-1 else '')); rr.font.name='Arial'; rr.font.size=Pt(6.5); rr.bold=(i==0); rr.italic=(filial_key(filial)=='CAMPO_GRANDE')
 linep=header.add_paragraph(); linep.paragraph_format.space_before=Pt(1); linep.paragraph_format.space_after=Pt(0); _set_bottom_border(linep)

def doc_footer(section):
    f=section.footer; p=f.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('ECOSUPPLY RECICLADORA LTDA  |  Contrato de Prestação de Serviços')
    r.font.name='Arial'; r.font.size=Pt(7)

def build_doc_v4(r,path):
 d=dict(r); f=filial_data(d.get('faturado_por')); doc=Document(); sec=doc.sections[0]
 sec.top_margin=Cm(2.95); sec.bottom_margin=Cm(1.35); sec.left_margin=Cm(1.55); sec.right_margin=Cm(1.55); sec.header_distance=Cm(.45); sec.footer_distance=Cm(.55)
 doc_header(sec,d.get('faturado_por')); doc_footer(sec)
 normal=doc.styles['Normal']; normal.font.name='Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); normal.font.size=Pt(9.1)
 normal.paragraph_format.space_after=Pt(4.0); normal.paragraph_format.line_spacing=1.0
 def p(text='',bold=False,center=False,after=4.0,before=0,keep=False,left=False):
  x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else (WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.JUSTIFY)
  x.paragraph_format.space_after=Pt(after); x.paragraph_format.space_before=Pt(before); x.paragraph_format.line_spacing=1.0; x.paragraph_format.keep_together=keep
  rr=x.add_run(text); rr.bold=bold; rr.font.name='Arial'; rr.font.size=Pt(9.1); return x
 def H(text): return p(text,True,False,3.5,2.5,True)
 # title aligned right as in the reference PDFs
 tbl=doc.add_table(rows=1,cols=2); tbl.autofit=False; tbl.columns[0].width=Cm(10.2); tbl.columns[1].width=Cm(7.2)
 for c in tbl.rows[0].cells: _set_cell_margins(c,0,0,0,0)
 tp=tbl.cell(0,1).paragraphs[0]; tp.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; tp.paragraph_format.space_after=Pt(5); tp.paragraph_format.line_spacing=.95
 tr=tp.add_run(f"INSTRUMENTO PARTICULAR DE CONTRATO DE PRESTAÇÃO DE SERVIÇOS DE RETIRADA, DECOMPOSIÇÃO, TRATAMENTO E DESCARTE, FIRMADO ENTRE {f['nome']} E {d.get('razao_social','').upper()}"); tr.bold=True; tr.font.name='Arial'; tr.font.size=Pt(8.7)
 H('IDENTIFICAÇÃO DAS PARTES CONTRATANTES')
 p(f"{d.get('razao_social','').upper()}, pessoa jurídica/pessoa física de direito privado com sede em {d.get('cidade','').upper()}, Estado de {d.get('uf','')}, na {d.get('endereco','')}, nº {d.get('numero','')} - Bairro {d.get('bairro','')} - CEP: {d.get('cep','')}, inscrita no CPF/CNPJ sob n.º {d.get('cpf_cnpj','')} e Inscrição Estadual n.º {d.get('ie','')}, neste ato representada por {d.get('responsavel','')}, portador(a) do RG n.º {d.get('rg','')} e CPF n.º {d.get('cpf_responsavel','')}, doravante denominada CONTRATANTE e")
 p(filial_text(d.get('faturado_por')))
 p('As partes acima identificadas têm, entre si, justo e acertado o presente Contrato de Prestação de Serviços, que se regerá pelas cláusulas seguintes e pelas condições de preço, forma e termo de pagamento descritas no presente.')
 H('DO OBJETO DO CONTRATO'); residuos=d.get('observacoes') or 'resíduos sólidos classe I e classe II conforme cadastro e caracterização fornecida pela CONTRATANTE'
 p('Cláusula 1ª. É objeto do presente contrato a prestação dos serviços de;')
 p(f'A) Serviço de coleta, triagem e envio para co-processamento de resíduos sólidos classe I e classe II, em especial: {residuos}')
 p('B) Fornecimento a título de comodato, sem custo adicional, de:',after=1.5)
 kits=[]
 for i in range(1,6):
  if d.get(f'kit_emb_{i}') and d.get(f'kit_qtd_{i}'):
   cor=f" - cor {d.get(f'kit_cor_{i}')}" if d.get(f'kit_cor_{i}') else ''
   kits.append(f"- {d[f'kit_qtd_{i}']} unidade(s) de {d[f'kit_emb_{i}']}{cor};")
 p('\n'.join(kits) if kits else '- Não há kit de coleta/comodato informado.',left=True)
 p('C) Transporte dos resíduos em caminhão licenciado para transporte de carga perigosa e com motorista habilitado com MOPP.')
 p('D) Destinação ambientalmente adequada de todos os resíduos coletados.')
 p('Parágrafo único – Caberá à CONTRATANTE providenciar às suas expensas, quando se fizer necessário, a caracterização dos resíduos de que trata esta Cláusula, através de análises completas executadas em entidades ou laboratórios de reconhecida capacidade técnica, conforme técnicas e métodos estabelecidos pela Norma Técnica, Legal ou Regulamentar pertinente, fornecendo, à CONTRATADA, cópia do laudo assinado por profissional competente.')
 H('OBRIGAÇÕES DA CONTRATANTE')
 p('Cláusula 2ª. A CONTRATANTE deverá abster-se de transferir para os locais de coleta quaisquer dos itens enumerados na cláusula primeira originados de outros estabelecimentos geradores, ainda que da mesma empresa ou grupo econômico.')
 p('Parágrafo Primeiro – Para fins deste instrumento, consideram-se “geradores” as pessoas físicas ou jurídicas que, em decorrência de sua atividade, gerem óleos lubrificantes usados ou contaminados, resíduos oleosos, estopas, embalagens, filtros de óleo e EPI´s contaminados com derivados de petróleo.')
 mtr=(d.get('emite_mtr') or '').upper()
 if 'CLIENTE' in mtr or 'CONTRATANTE' in mtr: p('Parágrafo Segundo – Compete ao CONTRATANTE/GERADOR a responsabilidade pela emissão do Manifesto de Transporte de Resíduos – MTR, por meio do SINIR – Sistema Nacional de Informações sobre a Gestão dos Resíduos Sólidos, para cada remessa de resíduos destinada à coleta e/ou destinação, nos termos da legislação vigente.')
 else: p('Parágrafo Segundo – Caso o CONTRATANTE/GERADOR opte pela operacionalização da emissão dos MTRs pela CONTRATADA, deverá cadastrá-la no SINIR com perfil de Administrador. O perfil de Administrador e as respectivas permissões serão utilizados pela CONTRATADA EXCLUSIVAMENTE para as atividades relacionadas à operacionalização de emissão dos MTRs das coletas realizadas pela ECOSUPPLY, mantendo com obrigatoriedade da CONTRATANTE todas as demais obrigações ambientais vinculadas a geração de resíduos e envio de declarações aos órgãos reguladores.')
 p('Cláusula 3ª. A CONTRATANTE deverá fornecer condições mínimas para atender a Segurança do Trabalho ao colaborador da Supply no momento da efetivação dos serviços objetos do presente instrumento, compreendendo todas as Normas Regulamentadoras do Ministério do Trabalho e pela legislação aplicável.

Parágrafo Único: A CONTRATANTE deverá atender, integralmente, a Norma Regulamentadora de n° 16, estando condicionada a prestação dos serviços ao fato de que toda a área de operação deverá abranger, no mínimo, círculo de distância do colaborador com raio de 10 metros com centro no ponto de abastecimento.')
 p('Cláusula 4ª. A CONTRATANTE deverá armazenar os itens discriminados na cláusula primeira em instalações adequadas devidamente licenciadas pelo órgão ambiental competente, em lugar acessível à coleta, utilizando os recipientes fornecidos a título de comodato pela CONTRATADA, de modo a não contaminar o meio ambiente.')
 H('OBRIGAÇÕES DA CONTRATADA')
 p('Cláusula 5ª. É dever da CONTRATADA observar todas as normas ambientais e técnicas pertinentes a execução dos serviços, dando plena e total garantia dos mesmos e responsabilizando-se isoladamente pela qualidade dos serviços prestados, devendo reparar ou refazer qualquer serviço que porventura for executado em desconformidade com as especificações, instruções e normas técnicas.')
 p('Cláusula 6ª. É dever da CONTRATADA estar devidamente autorizada pelo órgão regulador da indústria do petróleo e licenciada pelo órgão ambiental competente para realizar atividade de coleta e destinação de óleo lubrificante usado ou contaminado, limpeza de caixa separadora, coleta de resíduos oleosos, estopas, embalagens, filtros de óleo e EPI´S usados contaminados com derivados de petróleo.')
 per=d.get('periodicidade') or 'CONFORME PROGRAMAÇÃO'; fran=d.get('franquia_solidos') or '0'
 p(f'Cláusula 7ª. A CONTRATADA deverá efetuar a coleta e destinação na periodicidade {per}, em quantidade não superior a {fran} kg de todo material a ser coletado ({residuos}), sendo que a eventual necessidade de coleta fora da periodicidade programada ou em quantidade superior, implicará no pagamento de taxa adicional, cujo valor será fixado de acordo com o serviço a ser executado.')
 for n,t in [(8,'A CONTRATADA se reserva ao direito de devolver à CONTRATANTE, sem ônus para a primeira, todo o resíduo recebido que não possa processá-lo em razão de impedimento legal ou por inviabilidade técnica ou econômica.'),(9,'A CONTRATADA deverá manter em seu escritório de fábrica, um “livro de ocorrências” onde deverão ser escriturados todos os fatos que não se enquadrem como rotina.'),(10,'A CONTRATADA deverá comunicar à CONTRATANTE, por escrito ou por meio eletrônico, quaisquer ocorrências que impeçam o cumprimento das obrigações por ela aqui assumidas, indicando suas causas, ou efeitos e sugestões que devam ser tomadas.'),(11,'As partes “de per si”, se obrigam a obedecer as exigências, condições e determinações regulares do Órgão de Controle da Poluição Ambiental competente, responsabilizando-se, cada uma, isoladamente, pelas infrações que eventualmente cometerem.')]: p(f'Cláusula {n}ª. {t}')
 p('Cláusula 12ª. Para a execução dos serviços, a CONTRATADA disponibilizará os seguintes recursos:\nA) 01 caminhão Baú ou Rollon devidamente certificado e licenciado para transporte de Produtos Perigosos (INMETRO), condutores habilitados e treinados (MOPP) e contrato com empresa de atendimento emergencial, além de seguro ambiental da carga transportada;\nB) Mão de obra disponibilizada: Operadores habilitados e treinados (MOPP, NR-33, NR-35);\nC) EPI’s necessários para a realização dos serviços;\nD) Documentação: Check-List e Certificado de Coleta de Resíduos.')
 H('DA COLETA E DO TRANSPORTE')
 p('Cláusula 13ª. A CONTRATADA irá realizar as coletas no período diurno, no estabelecimento da CONTRATANTE, não sendo executadas aos sábados, domingos e feriados. Caso o dia programado para coleta for feriado, a mesma será realizada no dia útil anterior ou posterior àquela data.')
 p('Cláusula 14ª. A CONTRATANTE se responsabilizará de acompanhar a coleta, a pesagem, e o preenchimento do manifesto, assinando e dando ciência de todos os resíduos que foram coletados.')
 p('Cláusula 15ª. Ocorrendo impossibilidade real da CONTRATADA na execução do serviço, esta deverá ser reprogramada e comunicada imediatamente à CONTRATANTE para a sua realização. Neste caso, não será cobrado da CONTRATANTE valores correspondentes às quantidades excedentes e se não for possível realizar a coleta dentro do período, também não será cobrada a mensalidade correspondente.')
 H('DO PREÇO E DAS CONDIÇÕES DE PAGAMENTO')
 p('Cláusula 16ª. A CONTRATANTE pagará à CONTRATADA, mediante boleto bancário a título de remuneração pelos serviços avençados por este instrumento a importância de:')
 def moeda(v):
  z=str(v or '').strip(); return ('R$ '+z.replace('R$','').strip()) if z else ''
 if d.get('valor_faturamento'): p(f"- {moeda(d.get('valor_faturamento'))} referente ao serviço na periodicidade {per}.")
 for key,label in [('valor_excedente','por kg de material coletado excedente ao contratado'),('valor_terra','por kg de terra, serragem, graxa contaminados com derivados de petróleo'),('valor_agua','por litro de água'),('valor_lampada','por unidade de lâmpada')]:
  if d.get(key): p(f"- {moeda(d[key])} {label}, sendo que esse resíduo não faz parte da franquia mencionada na cláusula 7ª deste instrumento.")
 if d.get('outros_valores'): p(f"- Outros valores: {d['outros_valores']}.")
 venc=d.get('prazo_dia_fixo') or d.get('dia_limite_faturamento') or ''
 p(f"Devendo esses valores serem faturados conforme as condições comerciais cadastradas{(' e com vencimento no dia '+venc) if venc else ''}.")
 p('Parágrafo Primeiro: O atraso no pagamento de toda e qualquer fatura obrigará a Contratante no pagamento do principal acrescido de multa contratual de 2% (dois por cento) após atraso, além de juros de 1,0% (um por cento) ao mês, sobre o total do débito.\nI - O percentual estabelecido neste Parágrafo, a título de multa, poderá ser reavaliado em caso de alteração significativa da taxa de inflação mensal que se verifica nesta data.\nII - Em caso de cobrança judicial, devem ser acrescidas custas processuais e 20% de honorários advocatícios.')
 p('Parágrafo Segundo: A remuneração de que trata o “caput” desta Cláusula inclui o tratamento e a disposição final do efluente líquido e do resíduo sólido resultantes do processamento do resíduo de que trata a Cláusula Primeira deste Contrato, de modo a atender as exigências dos órgãos de meio ambiente competentes.')
 p('Parágrafo Terceiro: Eventuais tributos e contribuições de qualquer natureza que porventura venham a ser criadas ou elevadas serão objeto de repasse para os custos.')
 p(f"Parágrafo Quarto: A CONTRATADA emitirá Nota Fiscal Eletrônica atendendo determinação da legislação tributária e a remeterá para a CONTRATANTE no e-mail: {d.get('email_nfe') or '[E-MAIL NF-e]'}, o qual será instrumento hábil e regular para o recebimento do documento fiscal, sendo que a comprovação de entrega da nota se dará pela simples remessa do e-mail.")
 p('Parágrafo Quinto: Se por algum motivo o boleto bancário não puder ser pago e a CONTRATADA autorizar o pagamento por depósito bancário, deverá ser acrescido no valor do depósito bancário o valor correspondente à baixa bancária de título.')
 p('Parágrafo Sexto: O valor da remuneração estabelecido no “caput” será reajustado automaticamente a cada 12 (doze) meses através da aplicação do índice IGP-M ou outro índice ou reajuste pertinente.')
 p('Cláusula 17ª. No que se refere a incidência do Imposto Sobre Serviços - ISS, a CONTRATANTE deverá observar rigorosamente os dispositivos legais que tratam da incidência e eventual retenção, sendo que a exigência ou retenção indevida do tributo implicará no repasse do valor para o preço do serviço.')
 p('Cláusula 18ª. No caso de descumprimento de qualquer das cláusulas deste instrumento, a parte que não cumpriu deverá pagar multa correspondente ao valor dos serviços nos últimos três meses anteriores a infração contratual, além de ensejar a rescisão contratual motivada.')
 p('Cláusula 19ª. O descumprimento de qualquer das cláusulas deste instrumento permite à parte lesada a rescisão contratual motivada, bastando para tanto a notificação da parte contrária, que produzirá efeitos imediatos.')
 H('DA RESCISÃO IMOTIVADA')
 p('Cláusula 20ª. Poderá o presente instrumento ser rescindido por qualquer uma das partes, em qualquer momento, sem que haja qualquer tipo de motivo relevante, não obstante a outra parte deverá ser avisada previamente por escrito, no prazo de 30 (trinta) dias.')
 p('Cláusula 21ª. Caso a CONTRATANTE já tenha realizado o pagamento pelo serviço, e mesmo assim, requisite a rescisão imotivada do presente contrato, terá o valor da quantia paga devolvido, deduzindo-se 10% de taxas administrativas e impostos recolhidos.')
 p('Cláusula 22ª. Caso seja a CONTRATADA quem requeira a rescisão imotivada, deverá devolver a quantia que se refere aos serviços por ele não prestados a CONTRATANTE, acrescentado de 2% de taxas administrativas.')
 H('DO PRAZO'); p('Cláusula 23ª. O presente contrato é por prazo indeterminado iniciando-se no momento e na data da assinatura do mesmo.')
 H('DAS CONDIÇÕES GERAIS'); p('Cláusula 24ª. Salvo com a expressa autorização da CONTRATANTE, não pode a CONTRATADA transferir ou subcontratar os serviços previstos neste instrumento, sob o risco de ocorrer a rescisão imediata.'); p('Cláusula 25ª. A CONTRATADA emitirá o Certificado de Tratamento e Disposição final de Resíduos, no qual constará a quantidade tratada no período correspondente, sendo enviado ao cliente no prazo de 90 dias.')
 H('DO FORO'); p(f"Cláusula 26ª. Para dirimir quaisquer controvérsias oriundas do presente contrato, as partes elegem o foro da comarca de {f['foro']}.")
 p('Por estarem assim justos e contratados, firmam o presente instrumento em 02 (duas) vias de igual teor, juntamente com 02 (duas) testemunhas.')
 meses=['janeiro','fevereiro','março','abril','maio','junho','julho','agosto','setembro','outubro','novembro','dezembro']; now=datetime.now(); p(f"{f['cidade']}-{f['uf']}, {now.day} de {meses[now.month-1]} de {now.year}.")
 sig=doc.add_table(rows=1,cols=2); sig.autofit=False; sig.columns[0].width=Cm(8.7); sig.columns[1].width=Cm(8.7)
 for idx,(name,role) in enumerate([(f['nome'],'CONTRATADA'),(d.get('razao_social','').upper(),'CONTRATANTE')]):
  cc=sig.cell(0,idx); _set_cell_margins(cc,0,40,0,40); pp=cc.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=pp.add_run('\n\n________________________________\n'+role+'\n'+name); rr.bold=True; rr.font.name='Arial'; rr.font.size=Pt(8.4)
 doc.save(path)

def build_pdf_v4(r,path):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    d=dict(r); f=filial_data(d.get('faturado_por')); styles=getSampleStyleSheet()
    body=ParagraphStyle('Body',parent=styles['Normal'],fontName='Helvetica',fontSize=8.5,leading=11,alignment=TA_JUSTIFY,spaceAfter=5)
    head=ParagraphStyle('Head',parent=body,fontName='Helvetica-Bold',spaceBefore=4,spaceAfter=5)
    center=ParagraphStyle('Center',parent=body,alignment=TA_CENTER,fontName='Helvetica-Bold')
    def esc(s):
        import html; return html.escape(str(s or '')).replace('\n','<br/>')
    def header_footer(canvas,doc):
        canvas.saveState(); w,h=A4
        logo=BASE/'static'/'logo_ecosupply.png'
        if logo.exists(): canvas.drawImage(str(logo),1.7*cm,h-2.4*cm,width=6.8*cm,height=1.5*cm,preserveAspectRatio=True,mask='auto')
        canvas.setFont('Helvetica-Bold',7); canvas.drawString(11.2*cm,h-1.25*cm,f['cabecalho'][0])
        canvas.setFont('Helvetica',6.2); yy=h-1.55*cm
        for line in f['cabecalho'][1:4]: canvas.drawString(11.2*cm,yy,line); yy-=0.22*cm
        canvas.setLineWidth(1.5); canvas.line(1.7*cm,h-2.7*cm,w-1.7*cm,h-2.7*cm)
        canvas.setFont('Helvetica',6.5); canvas.drawCentredString(w/2,.8*cm,f"{f['nome']}  |  Contrato de Prestação de Serviços  |  Página {doc.page}")
        canvas.restoreState()
    pdf=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=1.7*cm,leftMargin=1.7*cm,topMargin=3.1*cm,bottomMargin=1.3*cm)
    story=[]
    title=f"INSTRUMENTO PARTICULAR DE CONTRATO DE PRESTAÇÃO DE SERVIÇOS DE RETIRADA, DECOMPOSIÇÃO, TRATAMENTO E DESCARTE, FIRMADO ENTRE {f['nome']} E {esc(d.get('razao_social','').upper())}"
    story.append(Table([['',Paragraph(title,head)]],colWidths=[10.7*cm,6.2*cm],style=[('VALIGN',(0,0),(-1,-1),'TOP')]))
    def H(x): story.append(Paragraph(esc(x),head))
    def P(x): story.append(Paragraph(esc(x),body))
    H('IDENTIFICAÇÃO DAS PARTES CONTRATANTES')
    P(f"{d.get('razao_social','').upper()}, pessoa jurídica/pessoa física de direito privado com sede em {d.get('cidade','').upper()}, Estado de {d.get('uf','')}, na {d.get('endereco','')}, nº {d.get('numero','')} - Bairro {d.get('bairro','')} - CEP: {d.get('cep','')}, inscrita no CPF/CNPJ sob n.º {d.get('cpf_cnpj','')} e Inscrição Estadual n.º {d.get('ie','')}, neste ato representada por {d.get('responsavel','')}, portador(a) do RG n.º {d.get('rg','')} e CPF n.º {d.get('cpf_responsavel','')}, doravante denominada CONTRATANTE e")
    P(filial_text(d.get('faturado_por'))); P('As partes acima identificadas têm, entre si, justo e acertado o presente Contrato de Prestação de Serviços, que se regerá pelas cláusulas seguintes e pelas condições de preço, forma e termo de pagamento descritas no presente.')
    H('DO OBJETO DO CONTRATO'); residuos=d.get('observacoes') or 'resíduos sólidos classe I e classe II conforme cadastro e caracterização fornecida pela CONTRATANTE'
    P(f"Cláusula 1ª. É objeto do presente contrato a prestação dos serviços de;\nA) Serviço de coleta, triagem e envio para co-processamento/tratamento de resíduos sólidos classe I e classe II, em especial: {residuos}\nB) Fornecimento a título de comodato, sem custo adicional, de:")
    kits=[]
    for i in range(1,6):
        if d.get(f'kit_emb_{i}') and d.get(f'kit_qtd_{i}'): kits.append(f"- {d[f'kit_qtd_{i}']} unidade(s) de {d[f'kit_emb_{i}']} - cor {d.get(f'kit_cor_{i}','')};")
    P('\n'.join(kits) if kits else '- Não há kit de coleta/comodato informado.')
    P('C) Transporte dos resíduos em caminhão licenciado para transporte de carga perigosa e com motorista habilitado com MOPP.\nD) Destinação ambientalmente adequada de todos os resíduos coletados.')
    P('Parágrafo único – Caberá à CONTRATANTE providenciar às suas expensas, quando se fizer necessário, a caracterização dos resíduos de que trata esta Cláusula, através de análises completas executadas em entidades ou laboratórios de reconhecida capacidade técnica, conforme técnicas e métodos estabelecidos pela Norma Técnica, Legal ou Regulamentar pertinente, fornecendo, à CONTRATADA, cópia do laudo assinado por profissional competente.')
    H('OBRIGAÇÕES DA CONTRATANTE'); P('Cláusula 2ª. A CONTRATANTE deverá abster-se de transferir para os locais de coleta quaisquer dos itens enumerados na cláusula primeira originados de outros estabelecimentos geradores, ainda que da mesma empresa ou grupo econômico.')
    P('Parágrafo Primeiro – Para fins deste instrumento, consideram-se “geradores” as pessoas físicas ou jurídicas que, em decorrência de sua atividade, gerem óleos lubrificantes usados ou contaminados, resíduos oleosos, estopas, embalagens, filtros de óleo e EPI´s contaminados com derivados de petróleo.')
    if 'CLIENTE' in (d.get('emite_mtr') or '').upper() or 'CONTRATANTE' in (d.get('emite_mtr') or '').upper(): P('Parágrafo Segundo – Compete ao CONTRATANTE/GERADOR a responsabilidade pela emissão do Manifesto de Transporte de Resíduos – MTR, por meio do SINIR – Sistema Nacional de Informações sobre a Gestão dos Resíduos Sólidos, para cada remessa de resíduos destinada à coleta e/ou destinação, nos termos da legislação vigente.')
    else: P('Parágrafo Segundo – Caso o CONTRATANTE/GERADOR opte pela operacionalização da emissão dos MTRs pela CONTRATADA, deverá cadastrá-la no SINIR com perfil e permissões necessárias, permanecendo com a CONTRATANTE as demais obrigações ambientais vinculadas à geração de resíduos e envio de declarações aos órgãos reguladores.')
    P('Cláusula 3ª. A CONTRATANTE deverá fornecer condições mínimas para atender a Segurança do Trabalho ao colaborador da Supply no momento da efetivação dos serviços objetos do presente instrumento, compreendendo todas as Normas Regulamentadoras do Ministério do Trabalho e pela legislação aplicável.

Parágrafo Único: A CONTRATANTE deverá atender, integralmente, a Norma Regulamentadora de n° 16, estando condicionada a prestação dos serviços ao fato de que toda a área de operação deverá abranger, no mínimo, círculo de distância do colaborador com raio de 10 metros com centro no ponto de abastecimento.'); P('Cláusula 4ª. A CONTRATANTE deverá armazenar os itens discriminados na cláusula primeira em instalações adequadas devidamente licenciadas pelo órgão ambiental competente, em lugar acessível à coleta, utilizando os recipientes fornecidos a título de comodato pela CONTRATADA, de modo a não contaminar o meio ambiente.')
    H('OBRIGAÇÕES DA CONTRATADA')
    texts=[(5,'É dever da CONTRATADA observar todas as normas ambientais e técnicas pertinentes a execução dos serviços, dando plena e total garantia dos mesmos e responsabilizando-se isoladamente pela qualidade dos serviços prestados, devendo reparar ou refazer qualquer serviço que porventura for executado em desconformidade com as especificações, instruções e normas técnicas.'),(6,'É dever da CONTRATADA estar devidamente autorizada pelo órgão regulador e licenciada pelo órgão ambiental competente para realizar as atividades objeto deste contrato.')]
    for n,t in texts:P(f'Cláusula {n}ª. {t}')
    per=d.get('periodicidade') or 'CONFORME PROGRAMAÇÃO'; fran=d.get('franquia_solidos') or '0'; P(f'Cláusula 7ª. A CONTRATADA deverá efetuar a coleta e destinação na periodicidade {per}, em quantidade não superior a {fran} kg de todo material a ser coletado ({residuos}), sendo que a eventual necessidade de coleta fora da periodicidade programada ou em quantidade superior implicará no pagamento de taxa adicional, cujo valor será fixado de acordo com o serviço a ser executado.')
    for n,t in [(8,'A CONTRATADA se reserva ao direito de devolver à CONTRATANTE, sem ônus para a primeira, todo o resíduo recebido que não possa processá-lo em razão de impedimento legal ou por inviabilidade técnica ou econômica.'),(9,'A CONTRATADA deverá manter em seu escritório de fábrica, um “livro de ocorrências” onde deverão ser escriturados todos os fatos que não se enquadrem como rotina.'),(10,'A CONTRATADA deverá comunicar à CONTRATANTE, por escrito ou por meio eletrônico, quaisquer ocorrências que impeçam o cumprimento das obrigações por ela aqui assumidas, indicando suas causas, ou efeitos e sugestões que devam ser tomadas.'),(11,'As partes “de per si”, se obrigam a obedecer as exigências, condições e determinações regulares do Órgão de Controle da Poluição Ambiental competente, responsabilizando-se, cada uma, isoladamente, pelas infrações que eventualmente cometerem.')]:P(f'Cláusula {n}ª. {t}')
    P('Cláusula 12ª. Para a execução dos serviços, a CONTRATADA disponibilizará os seguintes recursos:\nA) 01 caminhão Baú ou Rollon devidamente certificado e licenciado para transporte de Produtos Perigosos (INMETRO), condutores habilitados e treinados (MOPP) e contrato com empresa de atendimento emergencial, além de seguro ambiental da carga transportada;\nB) Mão de obra disponibilizada: Operadores habilitados e treinados (MOPP, NR-33, NR-35);\nC) EPI’s necessários para a realização dos serviços;\nD) Documentação: Check-List e Certificado de Coleta de Resíduos.')
    H('DA COLETA E DO TRANSPORTE'); P('Cláusula 13ª. A CONTRATADA irá realizar as coletas no período diurno, no estabelecimento da CONTRATANTE, não sendo executadas aos sábados, domingos e feriados. Caso o dia programado para coleta for feriado, a mesma será realizada no dia útil anterior ou posterior àquela data.'); P('Cláusula 14ª. A CONTRATANTE se responsabilizará de acompanhar a coleta, a pesagem, e o preenchimento do manifesto, assinando e dando ciência de todos os resíduos que foram coletados.'); P('Cláusula 15ª. Ocorrendo impossibilidade real da CONTRATADA na execução do serviço, esta deverá ser reprogramada e comunicada imediatamente à CONTRATANTE para a sua realização.')
    H('DO PREÇO E DAS CONDIÇÕES DE PAGAMENTO'); P('Cláusula 16ª. A CONTRATANTE pagará à CONTRATADA, mediante boleto bancário a título de remuneração pelos serviços avençados por este instrumento a importância de:')
    def moeda(v):
        s=str(v or '').strip(); return ('R$ '+s.replace('R$','').strip()) if s else ''
    if d.get('valor_faturamento'):P(f'- {moeda(d.get("valor_faturamento"))} referente ao serviço na periodicidade {per}.')
    for key,label in [('valor_excedente','por kg de material coletado excedente ao contratado'),('valor_terra','por kg de terra, serragem, graxa contaminados com derivados de petróleo'),('valor_agua','por litro de água'),('valor_lampada','por unidade de lâmpada')]:
        if d.get(key):P(f'- {moeda(d[key])} {label}, sendo que esse resíduo não faz parte da franquia mencionada na cláusula 7ª deste instrumento.')
    P('Parágrafo Primeiro: O atraso no pagamento de toda e qualquer fatura obrigará a Contratante no pagamento do principal acrescido de multa contratual de 2% (dois por cento) após atraso, além de juros de 1,0% (um por cento) ao mês, sobre o total do débito.\nI - O percentual estabelecido neste Parágrafo, a título de multa, poderá ser reavaliado em caso de alteração significativa da taxa de inflação mensal que se verifica nesta data.\nII - Em caso de cobrança judicial, devem ser acrescidas custas processuais e 20% de honorários advocatícios.')
    P('Parágrafo Segundo: A remuneração de que trata o “caput” desta Cláusula inclui o tratamento e a disposição final do efluente líquido e do resíduo sólido resultantes do processamento do resíduo de que trata a Cláusula Primeira deste Contrato, de modo a atender as exigências dos órgãos de meio ambiente competentes.'); P('Parágrafo Terceiro: Eventuais tributos e contribuições de qualquer natureza que porventura venham a ser criadas ou elevadas serão objeto de repasse para os custos.'); P(f"Parágrafo Quarto: A CONTRATADA emitirá Nota Fiscal Eletrônica atendendo determinação da legislação tributária e a remeterá para a CONTRATANTE no e-mail: {d.get('email_nfe') or '[E-MAIL NF-e]'}. A adoção de AntiSpam e demais meios eletrônicos que impeçam o recebimento do e-mail são de inteira responsabilidade da CONTRATANTE."); P('Parágrafo Quinto: Se por algum motivo o boleto bancário não puder ser pago e a CONTRATADA autorizar o pagamento por depósito bancário, deverá ser acrescido o valor referente à baixa bancária de título.'); P('Parágrafo Sexto: O valor da remuneração estabelecido no “caput” será reajustado automaticamente a cada 12 (doze) meses através da aplicação do índice IGP-M ou outro índice ou reajuste pertinente.')
    P('Cláusula 17ª. No que se refere a incidência do Imposto Sobre Serviços - ISS, a CONTRATANTE deverá observar rigorosamente os dispositivos legais que tratam da incidência e eventual retenção, sendo que a exigência ou retenção indevida do tributo implicará no repasse do valor para o preço do serviço.'); P('Cláusula 18ª. No caso de descumprimento de qualquer das cláusulas deste instrumento, a parte que não cumpriu deverá pagar multa correspondente ao valor dos serviços nos últimos três meses anteriores a infração contratual, além de ensejar a rescisão contratual motivada.'); P('Cláusula 19ª. O descumprimento de qualquer das cláusulas deste instrumento permite à parte lesada a rescisão contratual motivada, bastando para tanto a notificação da parte contrária, que produzirá efeitos imediatos.')
    H('DA RESCISÃO IMOTIVADA'); P('Cláusula 20ª. Poderá o presente instrumento ser rescindido por qualquer uma das partes, em qualquer momento, sem que haja qualquer tipo de motivo relevante, não obstante a outra parte deverá ser avisada previamente por escrito, no prazo de 30 (trinta) dias.'); P('Cláusula 21ª. Caso a CONTRATANTE já tenha realizado o pagamento pelo serviço, e mesmo assim, requisite a rescisão imotivada do presente contrato, terá o valor da quantia paga devolvido, deduzindo-se 10% de taxas administrativas e impostos recolhidos.'); P('Cláusula 22ª. Caso seja a CONTRATADA quem requeira a rescisão imotivada, deverá devolver a quantia que se refere aos serviços por ele não prestados a CONTRATANTE, acrescentado de 2% de taxas administrativas.')
    H('DO PRAZO'); P('Cláusula 23ª. O presente contrato é por prazo indeterminado iniciando-se no momento e na data da assinatura do mesmo.'); H('DAS CONDIÇÕES GERAIS'); P('Cláusula 24ª. Salvo com a expressa autorização da CONTRATANTE, não pode a CONTRATADA transferir ou subcontratar os serviços previstos neste instrumento, sob o risco de ocorrer a rescisão imediata.'); P('Cláusula 25ª. A CONTRATADA emitirá o Certificado de Tratamento e Disposição final de Resíduos, no qual constará a quantidade tratada no período correspondente, sendo enviado ao cliente no prazo de 90 dias.'); H('DO FORO'); P(f"Cláusula 26ª. Para dirimir quaisquer controvérsias oriundas do presente contrato, as partes elegem o foro da comarca de {f['foro']}."); P('Por estarem assim justos e contratados, firmam o presente instrumento em 02 (duas) vias de igual teor, juntamente com 02 (duas) testemunhas.')
    meses=['janeiro','fevereiro','março','abril','maio','junho','julho','agosto','setembro','outubro','novembro','dezembro']; now=datetime.now(); P(f"{f['cidade']}-{f['uf']}, {now.day} de {meses[now.month-1]} de {now.year}.")
    sig=Table([[Paragraph(f"<br/><br/>________________________________<br/><b>CONTRATADA</b><br/><b>{f['nome']}</b>",center),Paragraph(f'<br/><br/>________________________________<br/><b>CONTRATANTE</b><br/><b>{esc(d.get("razao_social","").upper())}</b>',center)]],colWidths=[8.4*cm,8.4*cm]); story.append(sig)
    pdf.build(story,onFirstPage=header_footer,onLaterPages=header_footer)

@app.route('/login',methods=['GET','POST'])
def login():
 if request.method=='POST':
  usuario=request.form.get('usuario','').strip(); senha=request.form.get('senha','')
  c=con(); u=c.execute('select * from usuarios where usuario=? and ativo=1',(usuario,)).fetchone(); c.close()
  if u and check_password_hash(u['senha_hash'],senha):
   session.clear(); session['user_id']=u['id']; session['usuario']=u['usuario']; session['nome']=u['nome']; session['perfil']=u['perfil']
   return redirect(request.args.get('next') or url_for('index'))
  flash('Usuário ou senha inválidos.')
 return render_template('login.html')

@app.route('/logout')
def logout():
 session.clear(); return redirect(url_for('login'))

@app.route('/usuarios',methods=['GET','POST'])
@admin_required
def usuarios():
 c=con()
 if request.method=='POST':
  usuario=request.form.get('usuario','').strip(); nome=request.form.get('nome','').strip(); senha=request.form.get('senha',''); perfil=request.form.get('perfil','USUARIO')
  if usuario and senha:
   try: c.execute('insert into usuarios(usuario,nome,senha_hash,perfil,ativo) values(?,?,?,?,1)',(usuario,nome,generate_password_hash(senha),perfil)); c.commit(); flash('Usuário criado.')
   except Exception: c.rollback(); flash('Não foi possível criar. Verifique se o usuário já existe.')
 rows=c.execute('select id,usuario,nome,perfil,ativo from usuarios order by usuario').fetchall(); c.close(); return render_template('usuarios.html',rows=rows)

@app.route('/acesso/<int:id>',methods=['GET','POST'])
@login_required
def acesso(id):
    c=con(); r=c.execute('select * from contratos where id=?',(id,)).fetchone(); c.close()
    if not r: return 'Contrato não encontrado',404
    if request.method=='POST':
        if request.form.get('senha')==ACTION_PASSWORD:
            ids=session.get('authorized_contracts',[])
            if id not in ids: ids.append(id)
            session['authorized_contracts']=ids
            return redirect(url_for('detalhe',id=id))
        flash('Senha incorreta.')
    return render_template('acesso.html',r=r)

@app.route('/status/<int:id>',methods=['POST'])
@login_required
def atualizar_status(id):
    guard=protect(id)
    if guard:return guard
    status=request.form.get('status','').strip()
    if status not in STATUS_OPTS: flash('Status inválido.'); return redirect(url_for('detalhe',id=id))
    c=con(); c.execute('update contratos set status=? where id=?',(status,id)); c.commit(); c.close(); flash('Status atualizado.'); return redirect(url_for('detalhe',id=id))

@app.route('/')
@login_required
def index():
 c=con(); rr=c.execute("select count(*) as total from contratos where status='PENDENTE'").fetchone(); total=rr['total']; c.close(); return render_template('index.html',total=total)
@app.route('/novo',methods=['GET','POST'])
@login_required
def novo():
 if request.method=='GET': return render_template('importar.html')
 f=request.files.get('arquivo');
 if not f: flash('Selecione a CAPA.xlsx'); return redirect(url_for('novo'))
 p=UP/f.filename; f.save(p)
 try:d,o=read_capa(p)
 except Exception as e: flash('Erro ao ler CAPA: '+str(e)); return redirect(url_for('novo'))
 return render_template('form.html',d=d,o=o,form_action=url_for('salvar'),button_text='Salvar como pendência',editing=False)
@app.route('/salvar',methods=['POST'])
@login_required
def salvar():
 d={k:request.form.get(k,'').strip() for k in FIELDS}; c=con(); rr=c.execute('select count(*) as total from contratos').fetchone(); n=rr['total']+1; chave=f'CT-{datetime.now().year}-{n:04d}'; cols=['chave','status','criado_em']+FIELDS; vals=[chave,'PENDENTE',datetime.now().isoformat(timespec='seconds')]+[d[k] for k in FIELDS]; c.execute(f"insert into contratos ({','.join(cols)}) values ({','.join(['?']*len(cols))})",vals); c.commit(); c.close(); flash(f'{chave} salvo em Pendências.'); return redirect(url_for('pendencias'))
@app.route('/pendencias')
@login_required
def pendencias():
 c=con(); rows=c.execute('select * from contratos order by id desc').fetchall(); c.close(); return render_template('pendencias.html',rows=rows)
@app.route('/contrato/<int:id>')
@login_required
def detalhe(id):
 guard=protect(id)
 if guard:return guard
 c=con(); r=c.execute('select * from contratos where id=?',(id,)).fetchone(); c.close(); return render_template('detalhe.html',r=r,status_opts=STATUS_OPTS)
@app.route('/editar/<int:id>',methods=['GET','POST'])
@login_required
def editar(id):
 guard=protect(id)
 if guard:return guard
 c=con(); r=c.execute('select * from contratos where id=?',(id,)).fetchone(); c.close()
 if not r:return 'Contrato não encontrado',404
 if request.method=='POST':
  d={k:request.form.get(k,'').strip() for k in FIELDS}; c=con(); sets=','.join([f'"{k}"=?' for k in FIELDS]); c.execute(f'update contratos set {sets} where id=?',[d[k] for k in FIELDS]+[id]); c.commit(); c.close(); flash('Cadastro atualizado com sucesso.'); return redirect(url_for('detalhe',id=id))
 # opções vêm do mesmo modelo oficial da CAPA
 try:
  _,o=read_capa(BASE/'CAPA_MODELO.xlsx')
 except Exception:
  o={'pagamento':['PRAZO DE PAGAMENTO','DIA PARA PAGAMENTO','DIA PARA PAGAMENTO FORA DO MÊS'],'filial':['ECONORTE','ECO CUIABA','ECO CAMPO GRANDE'],'representante':['FLÁVIO','VILA','MARCIA'],'mtr':['CLIENTE EMITE','ECO EMITE'],'cond':['SIM','NÃO'],'cor':[],'periodicidade':[],'embalagem':[],'categoria':[],'faturamento':[]}
 return render_template('form.html',d=dict(r),o=o,form_action=url_for('editar',id=id),button_text='Salvar alterações',editing=True)

@app.route('/gerar/word/<int:id>')
@login_required
def gerar_word(id):
 guard=protect(id)
 if guard:return guard
 c=con(); r=c.execute('select * from contratos where id=?',(id,)).fetchone(); c.close(); p=OUT/f"{r['chave']}_CONTRATO.docx"; build_doc_v4(r,p); c=con(); c.execute("update contratos set status='CONTRATO GERADO' where id=?",(id,)); c.commit(); c.close(); return send_file(p,as_attachment=True,download_name=p.name)
@app.route('/gerar/pdf/<int:id>')
@login_required
def gerar_pdf(id):
 guard=protect(id)
 if guard:return guard
 c=con(); r=c.execute('select * from contratos where id=?',(id,)).fetchone(); c.close(); pdf=OUT/f"{r['chave']}_CONTRATO.pdf"; build_pdf_v4(r,pdf); c=con(); c.execute("update contratos set status='CONTRATO GERADO' where id=?",(id,)); c.commit(); c.close(); return send_file(pdf,as_attachment=True,download_name=pdf.name,mimetype='application/pdf')
if __name__=='__main__': app.run(debug=False,host='0.0.0.0',port=int(os.getenv('PORT','5000')))
